#!/bin/sh
# Polyglot header: sh re-executes this file with the scripts/ project venv via uv,
# so the CLI works from any cwd. Python sees the next line as a string literal.
''''exec uv run --quiet --project "$(dirname "$(readlink -f "$0")")" python "$0" "$@" # '''
DESCRIPTION = """Convert a CSV file to XLSX, storing datetime columns as real Excel datetimes.

Examples:
  csv_to_xlsx.py fail_medisave.csv
  csv_to_xlsx.py in.csv -o out.xlsx
  csv_to_xlsx.py in.csv --datetime created_at updated_at msg_datetime
  csv_to_xlsx.py in.csv --no-auto-detect --datetime created_at
  csv_to_xlsx.py in.csv --tz-mode utc   # convert offset-aware values to UTC

Datetime columns are auto-detected (every non-empty value must parse) unless
--no-auto-detect is given. Columns named with --datetime are always treated as
datetimes; values that fail to parse are kept as text and reported.
"""
import argparse
import csv
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

FORMATS = (
    "%Y-%m-%d %H:%M:%S.%f",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y",
)


def parse_dt(s, tz_mode):
    """Parse a datetime string. Returns naive datetime or raises ValueError."""
    s = s.strip()
    # ISO 8601, possibly with offset, e.g. 2026-09-12T11:35:47+08:00
    try:
        d = datetime.fromisoformat(s.replace("Z", "+00:00"))
    except ValueError:
        d = None
        for fmt in FORMATS:
            try:
                d = datetime.strptime(s, fmt)
                break
            except ValueError:
                continue
        if d is None:
            raise ValueError(s)
    if d.tzinfo is not None:
        if tz_mode == "utc":
            d = d.astimezone(timezone.utc)
        d = d.replace(tzinfo=None)  # Excel has no timezone-aware cell type
    return d


def detect_datetime_columns(header, rows, tz_mode):
    found = []
    for i, name in enumerate(header):
        values = [r[i] for r in rows if i < len(r) and r[i].strip() != ""]
        if not values:
            continue
        try:
            for v in values:
                parse_dt(v, tz_mode)
        except ValueError:
            continue
        found.append(name)
    return found


def main(argv=None):
    ap = argparse.ArgumentParser(description=DESCRIPTION, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", type=Path, help="CSV file to convert")
    ap.add_argument("-o", "--output", type=Path, help="output .xlsx path (default: input with .xlsx)")
    ap.add_argument("--datetime", nargs="*", default=[], metavar="COL",
                    help="column names to treat as datetimes")
    ap.add_argument("--no-auto-detect", action="store_true",
                    help="only use columns from --datetime; do not auto-detect")
    ap.add_argument("--tz-mode", choices=["local", "utc"], default="local",
                    help="for values with a timezone offset: keep local wall-clock time (default) or convert to UTC")
    ap.add_argument("--number-format", default="yyyy-mm-dd hh:mm:ss",
                    help="Excel number format for datetime cells")
    ap.add_argument("--sheet", help="sheet name (default: input file stem)")
    ap.add_argument("--font", default="Arial", help="font name (default: Arial)")
    ap.add_argument("--delimiter", default=",", help="CSV delimiter (default: ,)")
    ap.add_argument("--encoding", default="utf-8-sig", help="CSV encoding (default: utf-8-sig)")
    ap.add_argument("--max-width", type=int, default=40, help="max column width (default: 40)")
    ap.add_argument("-q", "--quiet", action="store_true")
    args = ap.parse_args(argv)

    if not args.input.is_file():
        sys.exit(f"input not found: {args.input}")
    output = args.output or args.input.with_suffix(".xlsx")
    log = (lambda *a: None) if args.quiet else (lambda *a: print(*a, file=sys.stderr))

    with open(args.input, newline="", encoding=args.encoding) as f:
        reader = csv.reader(f, delimiter=args.delimiter)
        try:
            header = next(reader)
        except StopIteration:
            sys.exit("input CSV is empty")
        rows = list(reader)
    log(f"read {len(rows)} rows, {len(header)} columns")

    unknown = [c for c in args.datetime if c not in header]
    if unknown:
        sys.exit(f"--datetime column(s) not in header: {unknown}\nheader: {header}")

    dt_cols = list(args.datetime)
    if not args.no_auto_detect:
        for c in detect_datetime_columns(header, rows, args.tz_mode):
            if c not in dt_cols:
                dt_cols.append(c)
    dt_set = set(dt_cols)
    log(f"datetime columns: {dt_cols or 'none'}")

    parsed = {c: 0 for c in dt_cols}
    failed = []
    maxlen = [len(h) for h in header]
    out = []
    for rn, r in enumerate(rows, 2):
        if len(r) < len(header):
            r = r + [""] * (len(header) - len(r))
        elif len(r) > len(header):
            log(f"warning: row {rn} has {len(r)} fields, header has {len(header)}; extra fields dropped")
            r = r[:len(header)]
        new = []
        for i, (c, v) in enumerate(zip(header, r)):
            if c in dt_set:
                if v.strip() == "":
                    new.append(None)
                else:
                    try:
                        new.append(parse_dt(v, args.tz_mode))
                        parsed[c] += 1
                    except ValueError:
                        failed.append((rn, c, v))
                        new.append(v)
            else:
                maxlen[i] = max(maxlen[i], len(v))
                new.append(v if v != "" else None)
        out.append(new)

    for c in dt_cols:
        log(f"  {c}: {parsed[c]} parsed")
    if failed:
        log(f"warning: {len(failed)} datetime value(s) kept as text, e.g. {failed[:3]}")

    wb = Workbook()
    ws = wb.active
    ws.title = (args.sheet or args.input.stem)[:31]
    font = Font(name=args.font, size=10)
    bold = Font(name=args.font, size=10, bold=True)
    ws.append(header)
    for cell in ws[1]:
        cell.font = bold
    for row in out:
        ws.append(row)

    dt_idx = {i + 1 for i, c in enumerate(header) if c in dt_set}
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.font = font
            if cell.column in dt_idx:
                cell.number_format = args.number_format
    for i, c in enumerate(header, 1):
        width = max(len(args.number_format), len(c)) + 2 if c in dt_set else min(maxlen[i - 1] + 2, args.max_width)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    wb.save(output)
    log(f"saved {output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
